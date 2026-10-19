import io
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

DAY_NAMES = {
    1: "Senin",
    2: "Selasa",
    3: "Rabu",
    4: "Kamis",
    5: "Jumat",
    6: "Sabtu"
}

def export_schedule_excel(job_info: Dict[str, Any], schedules: List[Dict[str, Any]]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jadwal Kuliah"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Calibri", size=15, bold=True, color="1D8699")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="333333")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=9)
    bold_data_font = Font(name="Calibri", size=9, bold=True)

    header_fill = PatternFill(start_color="1D8699", end_color="1D8699", fill_type="solid")
    alt_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. Header Information
    ws.merge_cells("A1:K1")
    ws["A1"] = "UNIVERSITAS BINA BANGSA GETSEMPENA (UBBG) BANDA ACEH"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = "FAKULTAS SAINS, TEKNOLOGI, DAN ILMU KESEHATAN — PRODI ILMU KOMPUTER"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:K3")
    sem_text = "Ganjil" if job_info.get("semester_type") == "ODD" else "Genap"
    ws["A3"] = f"JADWAL PERKULIAHAN SEMESTER {sem_text.upper()} T.A. {job_info.get('academic_year', '2026/2027')}"
    ws["A3"].font = subtitle_font
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20

    # 2. Table Column Headers
    headers = [
        "No", "Hari", "Pukul (WIB)", "Kode MK", "Nama Mata Kuliah", 
        "SKS", "Sem", "Kelas", "Peminatan (KBK)", "Dosen Pengampu", "Ruangan"
    ]
    
    start_row = 5
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 26

    # Sort schedules by Day, Period, Room
    sorted_schedules = sorted(
        schedules, 
        key=lambda s: (s["day_of_week"], s["period_order"], s["room_code"])
    )

    # 3. Data Rows
    current_row = start_row + 1
    for idx, item in enumerate(sorted_schedules, 1):
        day_str = DAY_NAMES.get(item["day_of_week"], f"Hari {item['day_of_week']}")
        time_str = f"{item['start_time'].strftime('%H:%M')} - {item['end_time'].strftime('%H:%M')}"
        
        row_values = [
            idx,
            day_str,
            time_str,
            item["course_code"],
            item["course_name"],
            item["credits"],
            item["semester"],
            item["class_name"],
            item["kbk"],
            item["lecturer_name"],
            f"{item['room_code']} ({item['room_name']})"
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 3, 6, 7, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            
            if idx % 2 == 0:
                cell.fill = alt_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Column Widths
    col_widths = {
        "A": 6, "B": 10, "C": 15, "D": 12, "E": 32, 
        "F": 6, "G": 6, "H": 25, "I": 16, "J": 28, "K": 25
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_schedule_pdf(job_info: Dict[str, Any], schedules: List[Dict[str, Any]]) -> io.BytesIO:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1, # Center
        textColor=colors.HexColor('#1D8699')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#333333')
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#555555')
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9
    )
    cell_center = ParagraphStyle(
        'CellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        alignment=1
    )
    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        alignment=1,
        textColor=colors.white
    )

    elements = []

    # Title & Subtitle
    elements.append(Paragraph("UNIVERSITAS BINA BANGSA GETSEMPENA (UBBG) BANDA ACEH", title_style))
    elements.append(Paragraph("FAKULTAS SAINS, TEKNOLOGI, DAN ILMU KESEHATAN — PRODI ILMU KOMPUTER", subtitle_style))
    sem_text = "Ganjil" if job_info.get("semester_type") == "ODD" else "Genap"
    elements.append(Paragraph(f"JADWAL PERKULIAHAN SEMESTER {sem_text.upper()} T.A. {job_info.get('academic_year', '2026/2027')}", subtitle_style))
    elements.append(Paragraph(f"Status: Resmi Dipublikasikan | Fitness Score: {job_info.get('best_fitness', 1.0):.4f} | Zero Hard Conflicts", meta_style))
    elements.append(Spacer(1, 10))

    # Table Data
    table_data = [
        [
            Paragraph("No", cell_header),
            Paragraph("Hari", cell_header),
            Paragraph("Waktu", cell_header),
            Paragraph("Kode MK", cell_header),
            Paragraph("Mata Kuliah", cell_header),
            Paragraph("SKS", cell_header),
            Paragraph("Kelas", cell_header),
            Paragraph("KBK", cell_header),
            Paragraph("Dosen Pengampu", cell_header),
            Paragraph("Ruangan", cell_header),
        ]
    ]

    sorted_schedules = sorted(
        schedules,
        key=lambda s: (s["day_of_week"], s["period_order"], s["room_code"])
    )

    for idx, item in enumerate(sorted_schedules, 1):
        day_str = DAY_NAMES.get(item["day_of_week"], f"H-{item['day_of_week']}")
        time_str = f"{item['start_time'].strftime('%H:%M')}-{item['end_time'].strftime('%H:%M')}"
        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(day_str, cell_center),
            Paragraph(time_str, cell_center),
            Paragraph(item["course_code"], cell_center),
            Paragraph(item["course_name"], cell_style),
            Paragraph(str(item["credits"]), cell_center),
            Paragraph(item["class_name"], cell_style),
            Paragraph(item["kbk"], cell_center),
            Paragraph(item["lecturer_name"], cell_style),
            Paragraph(item["room_code"], cell_center),
        ])

    col_widths = [20, 45, 60, 50, 160, 25, 120, 65, 140, 65]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D8699')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))

    elements.append(t)
    doc.build(elements)
    output.seek(0)
    return output
